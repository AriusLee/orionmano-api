"""Counterparty Q&A on a valuation version (Eric 2026-10, item 3).

  1. create_batch   — questions pasted or uploaded; split into items, pinned
                      to the active valuation version.
  2. answer_batch   — background: Claude drafts a simplified and a detailed
                      answer per question, and when the honest answer is "the
                      model should move", proposes parameter changes. Each
                      change's impact on EV / equity is computed by the
                      deterministic model before anyone approves it.
  3. decide_change  — analyst approves / rejects (and may edit) each change.
  4. apply_item     — background: approved changes → a new valuation version
                      (trigger qa_change) + its written report, both linked to
                      the question. Changes are also pinned on the company so a
                      later full regenerate keeps what was agreed.
"""

from __future__ import annotations

import asyncio
import io
import logging
import re
import uuid
from datetime import datetime, timezone
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import async_session
from app.models.company import Company
from app.models.qa import QaBatch, QaItem, QaModelChange
from app.services.agent.skills.produce_valuation_inputs import PINNABLE_PARAMS
from app.services.valuation_model import model_io, versions

logger = logging.getLogger(__name__)

_ANSWER_CONCURRENCY = 3
_tasks: set[asyncio.Task] = set()


class QaError(ValueError):
    pass


def _spawn(coro) -> None:
    task = asyncio.create_task(coro)
    _tasks.add(task)
    task.add_done_callback(_tasks.discard)


# ─── 1. Questions in ─────────────────────────────────────────────────────────

_NUMBERED = re.compile(r"^\s*(?:Q(?:uestion)?\s*)?\(?\d{1,3}[\.\):\-]\s*", re.IGNORECASE)


def split_questions_simple(text: str) -> list[str]:
    """Numbered lines start new questions; unnumbered lines continue the
    previous one. Without any numbering, each non-empty paragraph is one."""
    lines = [ln.rstrip() for ln in (text or "").splitlines()]
    if any(_NUMBERED.match(ln) for ln in lines):
        out: list[str] = []
        for ln in lines:
            if not ln.strip():
                continue
            if _NUMBERED.match(ln):
                out.append(_NUMBERED.sub("", ln).strip())
            elif out:
                out[-1] = f"{out[-1]} {ln.strip()}"
        return [q for q in out if q]
    paras = re.split(r"\n\s*\n", text or "")
    return [re.sub(r"\s+", " ", p).strip() for p in paras if p.strip()]


async def split_questions(text: str) -> list[str]:
    """Pull the questions out of free text (an email, a table dump from a
    spreadsheet). Uses the LLM for messy input; plain numbered lists are split
    without it."""
    simple = split_questions_simple(text)
    looks_clean = simple and all(len(q) < 1200 for q in simple) and bool(
        re.search(r"^\s*(?:Q\s*)?\d{1,3}[\.\):]", text or "", re.MULTILINE)
    )
    if looks_clean:
        return simple
    try:
        from app.services.ai.client import generate_text
        raw = await generate_text(
            system_prompt=(
                "You extract counterparty due-diligence questions about a company valuation from a document. "
                "Return ONLY a JSON object {\"questions\": [\"...\", ...]} — each question verbatim (fix only "
                "line-break artefacts), in document order, one entry per distinct question. Skip greetings, "
                "signatures, headings and anything that is not a question or request for information."
            ),
            user_prompt=(text or "")[:60_000],
            max_tokens=6000,
            skill="qa_split",
        )
        parsed = model_io._parse_json_response(raw)
        qs = parsed.get("questions") if isinstance(parsed, dict) else None
        if isinstance(qs, list):
            qs = [str(q).strip() for q in qs if str(q).strip()]
            if qs:
                return qs
    except Exception as e:
        logger.warning("LLM question split failed, using line split: %s", e)
    return simple


def text_from_upload(filename: str, data: bytes) -> str:
    """Text of an uploaded question list (.xlsx, .docx, .pdf, .txt/.csv/.md)."""
    import os
    import tempfile
    from app.services.ai import document_parser as dp

    lower = (filename or "").lower()
    if lower.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="ignore")
    suffix = os.path.splitext(lower)[1] or ".bin"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(data)
        path = tmp.name
    try:
        if lower.endswith(".pdf"):
            return dp.extract_text_from_pdf(path)
        if lower.endswith(".docx"):
            return dp.extract_text_from_docx(path)
        if lower.endswith((".xlsx", ".xlsm")):
            return dp.extract_text_from_xlsx(path)
        raise QaError("Upload a .xlsx, .docx, .pdf or .txt file")
    finally:
        try:
            os.unlink(path)
        except OSError:
            pass


async def create_batch(
    db: AsyncSession, company_id: uuid.UUID, *, title: str, questions: list[str],
    source_filename: str | None, user_id: uuid.UUID | None,
) -> QaBatch:
    questions = [q.strip() for q in questions if q and q.strip()]
    if not questions:
        raise QaError("No questions found")
    base = await versions.active_version(db, company_id)
    if base is None:
        raise QaError("Generate a valuation workpaper first — questions are answered against it")
    batch = QaBatch(
        company_id=company_id,
        title=title.strip() or f"Questions {datetime.now(timezone.utc):%d %b %Y}",
        source_filename=source_filename,
        base_version_id=base.id,
        status="draft",
        created_by=user_id,
        items=[],
    )
    for i, q in enumerate(questions, start=1):
        batch.items.append(QaItem(company_id=company_id, item_no=i, question=q, status="pending", changes=[]))
    db.add(batch)
    await db.commit()
    await db.refresh(batch)
    return batch


# ─── 2. Drafting answers ─────────────────────────────────────────────────────

_SYSTEM = """You are a senior valuation partner at Orionmano Assurance Services. A counterparty (investor, underwriter or acquirer) has sent questions about the valuation model below. Draft the firm's response to ONE question.

Answer from the model and the company data only. Never invent a figure; if the model or data cannot answer something, say what is missing. Be direct and professional — this goes to the counterparty after the analyst reviews it.

Write two answers:
- "simplified_answer": 2-4 sentences for a non-specialist.
- "detailed_answer": the full response in markdown — the reasoning, the relevant figures from the model, the basis of each assumption touched, and (if a change is proposed) what changes and why. No inline citation numbers.

Decide whether the question exposes something the model SHOULD change (an assumption that is unsupported, inconsistent with the data, or that the counterparty has a fair point on). Only then propose changes, and only to the parameters listed as changeable, using their exact keys. Values are raw: percentages as fractions (0.025 = 2.5%), money in the model's currency and unit, exactly like the [raw …] values shown. If the model is defensible as it stands, propose no change and explain why in the answer.

Return ONLY a JSON object:
{
  "topic": "short topic, e.g. WACC, Revenue growth, Comparables, Historical financials",
  "simplified_answer": "...",
  "detailed_answer": "...",
  "changes": [{"param_key": "specific_risk_premium_pm", "new_value": 0.02, "rationale": "one or two sentences"}]
}"""


async def _company_context(db: AsyncSession, company_id: uuid.UUID) -> str:
    from app.services.agent.context import AgentContext
    ctx = AgentContext(db=db, company_id=company_id)
    await ctx.load_company_data()
    text = await ctx.get_best_context_str()
    return text[:40_000]


async def _draft_one(item_id: uuid.UUID, model_text: str, company_text: str, base_inputs: dict) -> None:
    async with async_session() as db:
        item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one_or_none()
        if item is None:
            return
        question = item.question
    user = (
        f"# Valuation model\n{model_text}\n\n# Company data\n{company_text}\n\n"
        f"# Question\n{question}"
    )
    try:
        parsed, _ = await model_io.ask_claude_json(_SYSTEM, user, max_tokens=4000)
        if not isinstance(parsed, dict):
            raise ValueError("the model did not return a JSON answer")
    except Exception as e:
        async with async_session() as db:
            item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one()
            item.status = "failed"
            item.error = f"{type(e).__name__}: {str(e)[:300]}"
            await db.commit()
        return

    proposals: list[tuple[str, float, str | None]] = []
    for ch in parsed.get("changes") or []:
        if not isinstance(ch, dict):
            continue
        key = str(ch.get("param_key") or "").strip()
        if key not in PINNABLE_PARAMS:
            continue
        try:
            new_value = float(ch.get("new_value"))
        except (TypeError, ValueError):
            continue
        if model_io.param_value(base_inputs, key) == new_value:
            continue
        proposals.append((key, new_value, str(ch.get("rationale") or "").strip() or None))

    async with async_session() as db:
        item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one()
        item.topic = str(parsed.get("topic") or "")[:80] or None
        item.simplified_answer = str(parsed.get("simplified_answer") or "").strip() or None
        item.detailed_answer = str(parsed.get("detailed_answer") or "").strip() or None
        item.error = None
        item.answered_at = datetime.now(timezone.utc)
        for key, new_value, rationale in proposals:
            item.changes.append(QaModelChange(
                param_key=key,
                path=PINNABLE_PARAMS[key][0],
                old_value=model_io.param_value(base_inputs, key),
                new_value=new_value,
                rationale=rationale,
                impact=model_io.impact(base_inputs, {key: new_value}),
                status="proposed",
            ))
        item.status = "change_proposed" if proposals else "answered"
        await db.commit()


async def answer_batch_bg(batch_id: uuid.UUID, item_ids: list[uuid.UUID] | None = None) -> None:
    async with async_session() as db:
        batch = (await db.execute(select(QaBatch).where(QaBatch.id == batch_id))).scalar_one_or_none()
        if batch is None:
            return
        base = await versions.get_version(db, batch.company_id, batch.base_version_id) if batch.base_version_id else None
        if base is None:
            base = await versions.active_version(db, batch.company_id)
        if base is None:
            batch.status = "failed"
            batch.progress_message = "No valuation version to answer against"
            await db.commit()
            return
        targets = [i for i in batch.items if (item_ids is None and i.status in ("pending", "failed")) or (item_ids and i.id in item_ids)]
        for i in targets:
            i.status = "pending"
            i.error = None
            if item_ids:
                # Redrafting replaces earlier, undecided proposals.
                for ch in list(i.changes):
                    if ch.status == "proposed":
                        i.changes.remove(ch)
        batch.status = "answering"
        batch.progress_message = f"Drafting {len(targets)} answer(s)…"
        await db.commit()
        model_text = model_io.describe_model(base.inputs, base.summary)
        company_text = await _company_context(db, batch.company_id)
        base_inputs = base.inputs
        ids = [i.id for i in targets]

    sem = asyncio.Semaphore(_ANSWER_CONCURRENCY)

    async def run(iid: uuid.UUID) -> None:
        async with sem:
            await _draft_one(iid, model_text, company_text, base_inputs)

    await asyncio.gather(*(run(i) for i in ids))

    async with async_session() as db:
        batch = (await db.execute(select(QaBatch).where(QaBatch.id == batch_id))).scalar_one()
        await db.refresh(batch, ["items"])
        failed = sum(1 for i in batch.items if i.status == "failed")
        batch.status = "answered"
        batch.progress_message = f"{failed} question(s) failed — retry them" if failed else None
        await db.commit()


def start_answering(batch_id: uuid.UUID, item_ids: list[uuid.UUID] | None = None) -> None:
    _spawn(answer_batch_bg(batch_id, item_ids))


# ─── 3. Analyst decisions ────────────────────────────────────────────────────

async def _item(db: AsyncSession, company_id: uuid.UUID, item_id: uuid.UUID) -> QaItem:
    item = (await db.execute(
        select(QaItem).where(QaItem.id == item_id, QaItem.company_id == company_id)
    )).scalar_one_or_none()
    if item is None:
        raise QaError("Question not found")
    return item


async def update_item(
    db: AsyncSession, company_id: uuid.UUID, item_id: uuid.UUID, *,
    simplified_answer: str | None = None, detailed_answer: str | None = None, question: str | None = None,
) -> QaItem:
    item = await _item(db, company_id, item_id)
    if question is not None and question.strip():
        item.question = question.strip()
    if simplified_answer is not None:
        item.simplified_answer = simplified_answer
    if detailed_answer is not None:
        item.detailed_answer = detailed_answer
    await db.commit()
    return item


async def decide_change(
    db: AsyncSession, company_id: uuid.UUID, change_id: uuid.UUID, *,
    approve: bool, new_value: float | None, user_id: uuid.UUID | None,
) -> QaModelChange:
    change = (await db.execute(select(QaModelChange).where(QaModelChange.id == change_id))).scalar_one_or_none()
    if change is None:
        raise QaError("Change not found")
    item = await _item(db, company_id, change.item_id)
    if item.status in ("applying", "applied"):
        raise QaError("This question's changes are already applied")
    if new_value is not None and new_value != change.new_value:
        base_inputs = await _base_inputs(db, item)
        change.new_value = new_value
        change.impact = model_io.impact(base_inputs, {change.param_key: new_value})
    change.status = "approved" if approve else "rejected"
    change.decided_by = user_id
    change.decided_at = datetime.now(timezone.utc)
    await db.commit()
    return change


async def add_change(
    db: AsyncSession, company_id: uuid.UUID, item_id: uuid.UUID, *, param_key: str, new_value: float, rationale: str | None,
) -> QaModelChange:
    """The analyst proposes a change the draft did not."""
    if param_key not in PINNABLE_PARAMS:
        raise QaError("That parameter cannot be changed from Q&A")
    item = await _item(db, company_id, item_id)
    if item.status in ("applying", "applied"):
        raise QaError("This question's changes are already applied")
    base_inputs = await _base_inputs(db, item)
    change = QaModelChange(
        item_id=item.id,
        param_key=param_key,
        path=PINNABLE_PARAMS[param_key][0],
        old_value=model_io.param_value(base_inputs, param_key),
        new_value=new_value,
        rationale=rationale,
        impact=model_io.impact(base_inputs, {param_key: new_value}),
        status="proposed",
    )
    db.add(change)
    if item.status == "answered":
        item.status = "change_proposed"
    await db.commit()
    return change


async def _base_inputs(db: AsyncSession, item: QaItem) -> dict:
    batch = (await db.execute(select(QaBatch).where(QaBatch.id == item.batch_id))).scalar_one()
    v = await versions.get_version(db, item.company_id, batch.base_version_id) if batch.base_version_id else None
    if v is None:
        v = await versions.active_version(db, item.company_id)
    return v.inputs if v is not None else {}


async def combined_impact(db: AsyncSession, company_id: uuid.UUID, item_id: uuid.UUID) -> dict[str, Any] | None:
    """Impact of this question's approved changes together, on the CURRENT
    active version (what apply_item will actually do)."""
    item = await _item(db, company_id, item_id)
    approved = {c.param_key: c.new_value for c in item.changes if c.status == "approved"}
    if not approved:
        return None
    active = await versions.active_version(db, company_id)
    if active is None:
        return None
    return model_io.impact(active.inputs, approved)


# ─── 4. Apply → new version + report ─────────────────────────────────────────

async def apply_item(db: AsyncSession, company_id: uuid.UUID, item_id: uuid.UUID, user_id: uuid.UUID, pin: bool = True) -> QaItem:
    item = await _item(db, company_id, item_id)
    approved = [c for c in item.changes if c.status == "approved"]
    if not approved:
        raise QaError("Approve at least one change first")
    if item.status in ("applying", "applied"):
        raise QaError("Already applied")
    item.status = "applying"
    item.error = None
    await db.commit()
    _spawn(_apply_bg(company_id, item.id, user_id, pin))
    return item


async def _apply_bg(company_id: uuid.UUID, item_id: uuid.UUID, user_id: uuid.UUID, pin: bool) -> None:
    from app.services.agent.context import AgentContext
    from app.services.agent.registry import registry
    from app.services.agent.skill import SkillStatus
    from app.services.agent.skills.valuation_inputs_schema import validate as validate_inputs

    async with async_session() as db:
        item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one()
        approved = {c.param_key: c.new_value for c in item.changes if c.status == "approved"}
        try:
            active = await versions.active_version(db, company_id)
            if active is None:
                raise QaError("No active valuation version")
            # Apply on top of the CURRENT model, so several answered
            # questions build on each other instead of reverting one another.
            new_inputs = model_io.patched(active.inputs, approved)
            model, error = validate_inputs(new_inputs)
            if model is None:
                raise QaError(f"Changed inputs failed validation: {error}")
            labels = "; ".join(
                f"{model_io.PARAM_LABELS.get(k, k)} {model_io.format_param(k, model_io.param_value(active.inputs, k))} → {model_io.format_param(k, v)}"
                for k, v in approved.items()
            )
            skill = registry.get("generate_valuation_workpaper")
            ctx = AgentContext(db=db, company_id=company_id, user_id=user_id)
            result = await skill.execute(
                ctx,
                inputs=new_inputs,
                version_trigger="qa_change",
                version_note=f"Q&A #{item.item_no}: {labels}",
                qa_item_id=item.id,
            )
            if result.status == SkillStatus.FAILED:
                raise QaError(result.message)
            data = result.data or {}
            item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one()
            item.resulting_version_id = uuid.UUID(data["version_id"]) if data.get("version_id") else None
            item.resulting_report_id = uuid.UUID(data["report_id"]) if data.get("report_id") else None
            item.status = "applied"
            if pin:
                company = (await db.execute(select(Company).where(Company.id == company_id))).scalar_one()
                pins = dict(company.pinned_overrides or {})
                pins.update(approved)
                company.pinned_overrides = pins
            await db.commit()
        except Exception as e:
            logger.exception("applying Q&A item %s failed", item_id)
            await db.rollback()
            item = (await db.execute(select(QaItem).where(QaItem.id == item_id))).scalar_one()
            item.status = "change_proposed"
            item.error = f"Apply failed: {str(e)[:400]}"
            await db.commit()


# ─── Read model + export ─────────────────────────────────────────────────────

def change_dict(c: QaModelChange, currency: str | None, unit: str | None) -> dict[str, Any]:
    return {
        "id": str(c.id),
        "param_key": c.param_key,
        "label": model_io.PARAM_LABELS.get(c.param_key, c.param_key),
        "kind": PINNABLE_PARAMS.get(c.param_key, ("", "number"))[1],
        "old_value": c.old_value,
        "new_value": c.new_value,
        "old_display": model_io.format_param(c.param_key, c.old_value, currency, unit),
        "new_display": model_io.format_param(c.param_key, c.new_value, currency, unit),
        "rationale": c.rationale,
        "impact": c.impact,
        "status": c.status,
    }


def item_dict(i: QaItem, currency: str | None, unit: str | None) -> dict[str, Any]:
    return {
        "id": str(i.id),
        "item_no": i.item_no,
        "question": i.question,
        "topic": i.topic,
        "simplified_answer": i.simplified_answer,
        "detailed_answer": i.detailed_answer,
        "status": i.status,
        "error": i.error,
        "resulting_version_id": str(i.resulting_version_id) if i.resulting_version_id else None,
        "resulting_report_id": str(i.resulting_report_id) if i.resulting_report_id else None,
        "answered_at": i.answered_at.isoformat() if i.answered_at else None,
        "changes": [change_dict(c, currency, unit) for c in i.changes],
    }


async def batch_dict(db: AsyncSession, batch: QaBatch, *, with_items: bool = True) -> dict[str, Any]:
    base = await versions.get_version(db, batch.company_id, batch.base_version_id) if batch.base_version_id else None
    currency = base.currency if base else None
    unit = base.unit if base else None
    out = {
        "id": str(batch.id),
        "title": batch.title,
        "source_filename": batch.source_filename,
        "status": batch.status,
        "progress_message": batch.progress_message,
        "base_version_id": str(batch.base_version_id) if batch.base_version_id else None,
        "base_version_no": base.version_no if base else None,
        "currency": currency,
        "unit": unit,
        "created_at": batch.created_at.isoformat() if batch.created_at else None,
        "item_count": len(batch.items),
        "answered_count": sum(1 for i in batch.items if i.status not in ("pending", "failed")),
        "applied_count": sum(1 for i in batch.items if i.status == "applied"),
    }
    if with_items:
        version_nos = {}
        ids = {i.resulting_version_id for i in batch.items if i.resulting_version_id}
        for vid in ids:
            v = await versions.get_version(db, batch.company_id, vid)
            if v is not None:
                version_nos[str(vid)] = v.version_no
        items = []
        for i in batch.items:
            d = item_dict(i, currency, unit)
            d["resulting_version_no"] = version_nos.get(d["resulting_version_id"])
            items.append(d)
        out["items"] = items
    return out


def _impact_text(c: dict) -> str:
    imp = c.get("impact") or {}
    pct = imp.get("ev_change_pct")
    return f"EV {pct * 100:+.1f}%" if isinstance(pct, (int, float)) else ""


def export_xlsx(batch: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Q&A"
    headers = ["No.", "Topic", "Question", "Simplified answer", "Detailed answer", "Model change", "Impact", "Status", "Resulting version"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for it in batch["items"]:
        changes = [c for c in it["changes"] if c["status"] != "rejected"]
        ws.append([
            it["item_no"],
            it.get("topic") or "",
            it["question"],
            it.get("simplified_answer") or "",
            it.get("detailed_answer") or "",
            "\n".join(f"{c['label']}: {c['old_display']} → {c['new_display']} ({c['status']})" for c in changes) or "No change",
            "\n".join(_impact_text(c) for c in changes if _impact_text(c)),
            it["status"].replace("_", " "),
            f"v{it['resulting_version_no']}" if it.get("resulting_version_no") else "",
        ])
    widths = [6, 18, 50, 50, 80, 45, 14, 14, 12]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx(batch: dict, company_name: str) -> bytes:
    from docx import Document as Docx
    from docx.shared import Pt

    doc = Docx()
    doc.add_heading(f"{company_name} — Responses to counterparty questions", level=1)
    doc.add_paragraph(f"{batch['title']} · answered against valuation v{batch.get('base_version_no') or '—'}")
    for it in batch["items"]:
        doc.add_heading(f"{it['item_no']}. {it['question']}", level=2)
        if it.get("simplified_answer"):
            p = doc.add_paragraph()
            p.add_run("In brief: ").bold = True
            p.add_run(it["simplified_answer"])
        if it.get("detailed_answer"):
            for para in re.split(r"\n\s*\n", it["detailed_answer"]):
                txt = para.strip()
                if not txt:
                    continue
                if txt.startswith("#"):
                    doc.add_heading(txt.lstrip("# ").strip(), level=3)
                else:
                    doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", txt))
        applied = [c for c in it["changes"] if c["status"] == "approved"]
        if applied:
            p = doc.add_paragraph()
            p.add_run("Model change: ").bold = True
            p.add_run("; ".join(f"{c['label']} {c['old_display']} → {c['new_display']}" for c in applied))
    for style in ("Normal",):
        doc.styles[style].font.size = Pt(10.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
